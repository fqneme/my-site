from __future__ import annotations

import hmac
import os
import re
import secrets
import sqlite3
from datetime import datetime
from io import BytesIO
from functools import wraps
from pathlib import Path

from flask import (
    Flask,
    abort,
    flash,
    g,
    redirect,
    render_template,
    request,
    session,
    send_file,
    url_for,
)
from werkzeug.security import check_password_hash, generate_password_hash
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DB_PATH = DATA_DIR / "site.db"

app = Flask(__name__)
# Для локального учебного проекта. Перед публикацией задайте FLASK_SECRET_KEY
# длинным случайным значением в переменных окружения.
app.config["SECRET_KEY"] = os.environ.get(
    "FLASK_SECRET_KEY", "local-development-key-change-before-deployment"
)
app.config["MAX_CONTENT_LENGTH"] = 2 * 1024 * 1024
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"

ALLOWED_DEVICES = {
    "Компьютер",
    "Ноутбук",
    "Телефон",
    "Планшет",
    "Другая техника",
}
STATUS_NAMES = {
    "Новая": "status-new",
    "В работе": "status-progress",
    "Ожидает оплаты": "status-payment",
    "Выполнена": "status-done",
    "Отменена": "status-cancelled",
}


def get_db_connection() -> sqlite3.Connection:
    """Открывает соединение с SQLite и включает поддержку внешних ключей."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def init_db() -> None:
    """Создаёт таблицы и безопасно дополняет базу из первой версии сайта."""
    with get_db_connection() as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                full_name TEXT NOT NULL,
                email TEXT NOT NULL UNIQUE COLLATE NOCASE,
                password_hash TEXT NOT NULL,
                created_at TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT 'user' CHECK (role IN ('user', 'admin'))
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS service_requests (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                customer_name TEXT NOT NULL,
                phone TEXT NOT NULL,
                device_type TEXT NOT NULL,
                problem_description TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'Новая',
                created_at TEXT NOT NULL,
                user_id INTEGER,
                FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE SET NULL
            )
            """
        )

        # В первой версии у заявок не было user_id. Добавляем поле, не удаляя
        # существующие заявки и не пересоздавая файл базы данных.
        columns = {
            row["name"]
            for row in connection.execute("PRAGMA table_info(service_requests)").fetchall()
        }
        if "user_id" not in columns:
            connection.execute(
                "ALTER TABLE service_requests ADD COLUMN user_id INTEGER REFERENCES users(id) ON DELETE SET NULL"
            )

        # Добавляем роль существующим аккаунтам, не удаляя данные.
        user_columns = {
            row["name"]
            for row in connection.execute("PRAGMA table_info(users)").fetchall()
        }
        if "role" not in user_columns:
            connection.execute(
                "ALTER TABLE users ADD COLUMN role TEXT NOT NULL DEFAULT 'user'"
            )


def get_current_user():
    user_id = session.get("user_id")
    if not user_id:
        return None
    with get_db_connection() as connection:
        return connection.execute(
            "SELECT id, full_name, email, created_at, role FROM users WHERE id = ?",
            (user_id,),
        ).fetchone()


def generate_csrf_token() -> str:
    token = session.get("csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["csrf_token"] = token
    return token


@app.before_request
def protect_post_requests():
    """Защищает формы от отправки с посторонних сайтов и загружает аккаунт."""
    if request.method == "POST":
        expected = session.get("csrf_token", "")
        submitted = request.form.get("csrf_token", "")
        if not expected or not submitted or not hmac.compare_digest(expected, submitted):
            flash("Форма устарела или не прошла проверку. Попробуй отправить её ещё раз.", "error")
            if request.endpoint in {"login", "register"}:
                return redirect(request.path)
            return redirect(url_for("index") + "#contact")

    g.user = get_current_user()


@app.context_processor
def inject_template_helpers():
    return {
        "current_user": getattr(g, "user", None),
        "csrf_token": generate_csrf_token,
        "status_classes": STATUS_NAMES,
    }


def login_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if getattr(g, "user", None) is None:
            flash("Сначала войди в аккаунт.", "error")
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapped_view


def admin_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        user = getattr(g, "user", None)
        if user is None:
            flash("Войди в аккаунт администратора.", "error")
            return redirect(url_for("login"))
        if user["role"] != "admin":
            flash("У этой страницы доступ только для администратора.", "error")
            return redirect(url_for("account"))
        return view(*args, **kwargs)

    return wrapped_view


def is_valid_email(email: str) -> bool:
    return re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email) is not None


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/contacts")
def contacts():
    """Отдельная страница контактов и порядка обращения в сервис."""
    return render_template("contacts.html")


@app.post("/request")
def create_request():
    customer_name = (request.form.get("customer_name") or "").strip()
    phone = (request.form.get("phone") or "").strip()
    device_type = (request.form.get("device_type") or "").strip()
    problem_description = (request.form.get("problem_description") or "").strip()

    errors = []
    if not 2 <= len(customer_name) <= 80:
        errors.append("Укажи имя длиной от 2 до 80 символов.")

    phone_digits = re.sub(r"\D", "", phone)
    phone_format_is_valid = re.fullmatch(r"\+?[0-9\s().-]+", phone) is not None
    if not phone_format_is_valid or not 7 <= len(phone_digits) <= 15:
        errors.append(
            "Проверь номер телефона: нужно от 7 до 15 цифр. "
            "Можно использовать +, пробелы, скобки, точку и дефис."
        )
    if device_type not in ALLOWED_DEVICES:
        errors.append("Выбери тип устройства из списка.")
    if not 5 <= len(problem_description) <= 1200:
        errors.append("Опиши проблему (от 5 до 1200 символов).")

    if errors:
        for error in errors:
            flash(error, "error")
        return redirect(url_for("index") + "#contact")

    created_at = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %z")
    current_user = getattr(g, "user", None)
    with get_db_connection() as connection:
        connection.execute(
            """
            INSERT INTO service_requests
                (customer_name, phone, device_type, problem_description, status, created_at, user_id)
            VALUES (?, ?, ?, ?, 'Новая', ?, ?)
            """,
            (
                customer_name,
                phone,
                device_type,
                problem_description,
                created_at,
                current_user["id"] if current_user else None,
            ),
        )

    flash("Заявка принята! Мы свяжемся с тобой для уточнения деталей.", "success")
    if current_user:
        return redirect(url_for("account"))
    return redirect(url_for("index") + "#contact")


@app.route("/register", methods=["GET", "POST"])
def register():
    if getattr(g, "user", None) is not None:
        return redirect(url_for("account"))

    if request.method == "GET":
        return render_template("register.html")

    full_name = (request.form.get("full_name") or "").strip()
    email = (request.form.get("email") or "").strip().lower()
    password = request.form.get("password") or ""
    password_confirm = request.form.get("password_confirm") or ""

    errors = []
    if not 2 <= len(full_name) <= 80:
        errors.append("Имя должно содержать от 2 до 80 символов.")
    if len(email) > 254 or not is_valid_email(email):
        errors.append("Введи корректный адрес электронной почты.")
    if len(password) < 8:
        errors.append("Пароль должен содержать минимум 8 символов.")
    if len(password) > 128:
        errors.append("Пароль слишком длинный: максимум 128 символов.")
    if password != password_confirm:
        errors.append("Пароли не совпадают.")

    if errors:
        for error in errors:
            flash(error, "error")
        return redirect(url_for("register"))

    created_at = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %z")
    try:
        with get_db_connection() as connection:
            cursor = connection.execute(
                """
                INSERT INTO users (full_name, email, password_hash, created_at)
                VALUES (?, ?, ?, ?)
                """,
                (full_name, email, generate_password_hash(password), created_at),
            )
            user_id = cursor.lastrowid
    except sqlite3.IntegrityError:
        flash("Аккаунт с такой почтой уже существует. Войди или используй другую почту.", "error")
        return redirect(url_for("register"))

    session.clear()
    session["user_id"] = user_id
    flash("Аккаунт создан! Добро пожаловать в Компьютер-Софт.", "success")
    return redirect(url_for("account"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if getattr(g, "user", None) is not None:
        return redirect(url_for("account"))

    if request.method == "GET":
        return render_template("login.html")

    email = (request.form.get("email") or "").strip().lower()
    password = request.form.get("password") or ""
    with get_db_connection() as connection:
        user = connection.execute(
            "SELECT id, password_hash FROM users WHERE email = ? COLLATE NOCASE",
            (email,),
        ).fetchone()

    if user is None or not check_password_hash(user["password_hash"], password):
        flash("Неверная почта или пароль.", "error")
        return redirect(url_for("login"))

    session.clear()
    session["user_id"] = user["id"]
    flash("Ты успешно вошёл в аккаунт.", "success")
    return redirect(url_for("account"))


@app.post("/logout")
def logout():
    session.clear()
    flash("Ты вышел из аккаунта.", "success")
    return redirect(url_for("index"))


@app.get("/account")
@login_required
def account():
    with get_db_connection() as connection:
        orders = connection.execute(
            """
            SELECT id, phone, device_type, problem_description, status, created_at
            FROM service_requests
            WHERE user_id = ?
            ORDER BY id DESC
            """,
            (g.user["id"],),
        ).fetchall()
    return render_template("account.html", orders=orders)


@app.get("/admin")
@admin_required
def admin_dashboard():
    """Панель администратора: статистика, поиск, фильтры и все заявки."""
    search = (request.args.get("q") or "").strip()[:100]
    selected_status = (request.args.get("status") or "Все").strip()
    allowed_statuses = tuple(STATUS_NAMES.keys())
    if selected_status not in (*allowed_statuses, "Все"):
        selected_status = "Все"

    conditions = []
    params = []
    if selected_status != "Все":
        conditions.append("sr.status = ?")
        params.append(selected_status)
    if search:
        conditions.append(
            "(CAST(sr.id AS TEXT) LIKE ? OR sr.customer_name LIKE ? OR sr.phone LIKE ? "
            "OR sr.device_type LIKE ? OR sr.problem_description LIKE ? OR u.email LIKE ?)"
        )
        search_value = f"%{search}%"
        params.extend([search_value] * 6)

    sql = """
        SELECT sr.id, sr.customer_name, sr.phone, sr.device_type,
               sr.problem_description, sr.status, sr.created_at,
               sr.user_id, u.email AS account_email
        FROM service_requests AS sr
        LEFT JOIN users AS u ON u.id = sr.user_id
    """
    if conditions:
        sql += " WHERE " + " AND ".join(conditions)
    sql += " ORDER BY sr.id DESC LIMIT 200"

    with get_db_connection() as connection:
        requests_list = connection.execute(sql, params).fetchall()
        user_count = connection.execute("SELECT COUNT(*) FROM users").fetchone()[0]
        request_count = connection.execute("SELECT COUNT(*) FROM service_requests").fetchone()[0]
        status_rows = connection.execute(
            "SELECT status, COUNT(*) AS amount FROM service_requests GROUP BY status"
        ).fetchall()

    status_counts = {status: 0 for status in allowed_statuses}
    for row in status_rows:
        status_counts[row["status"]] = row["amount"]

    return render_template(
        "admin.html",
        requests_list=requests_list,
        search=search,
        selected_status=selected_status,
        statuses=allowed_statuses,
        status_counts=status_counts,
        user_count=user_count,
        request_count=request_count,
        visible_count=len(requests_list),
    )


@app.post("/admin/requests/<int:request_id>/status")
@admin_required
def admin_update_request_status(request_id: int):
    new_status = (request.form.get("status") or "").strip()
    if new_status not in STATUS_NAMES:
        flash("Выбран неизвестный статус.", "error")
        return redirect(url_for("admin_dashboard"))

    search = (request.form.get("q") or "").strip()[:100]
    selected_status = (request.form.get("current_status") or "Все").strip()
    if selected_status not in (*STATUS_NAMES.keys(), "Все"):
        selected_status = "Все"

    with get_db_connection() as connection:
        cursor = connection.execute(
            "UPDATE service_requests SET status = ? WHERE id = ?",
            (new_status, request_id),
        )
        changed = cursor.rowcount

    if changed:
        flash(f"Статус заявки №{request_id} изменён: {new_status}.", "success")
    else:
        flash("Заявка не найдена.", "error")
    return redirect(url_for("admin_dashboard", status=selected_status, q=search))


@app.get("/admin/export.xlsx")
@admin_required
def admin_export_excel():
    """Создаёт Excel-отчёт по всем заявкам с учётом выбранных фильтров."""
    search = (request.args.get("q") or "").strip()[:100]
    selected_status = (request.args.get("status") or "Все").strip()
    allowed_statuses = tuple(STATUS_NAMES.keys())
    if selected_status not in (*allowed_statuses, "Все"):
        selected_status = "Все"

    conditions = []
    params = []
    if selected_status != "Все":
        conditions.append("sr.status = ?")
        params.append(selected_status)
    if search:
        conditions.append(
            "(CAST(sr.id AS TEXT) LIKE ? OR sr.customer_name LIKE ? OR sr.phone LIKE ? "
            "OR sr.device_type LIKE ? OR sr.problem_description LIKE ? OR u.email LIKE ?)"
        )
        search_value = f"%{search}%"
        params.extend([search_value] * 6)

    sql = """
        SELECT sr.id, sr.customer_name, sr.phone, sr.device_type,
               sr.problem_description, sr.status, sr.created_at, u.email AS account_email
        FROM service_requests AS sr
        LEFT JOIN users AS u ON u.id = sr.user_id
    """
    if conditions:
        sql += " WHERE " + " AND ".join(conditions)
    sql += " ORDER BY sr.id DESC"

    with get_db_connection() as connection:
        rows = connection.execute(sql, params).fetchall()
        user_count = connection.execute("SELECT COUNT(*) FROM users").fetchone()[0]

    # Апостроф перед опасными первыми символами не даёт Excel трактовать
    # пользовательский текст как формулу при открытии отчёта.
    def safe_excel_value(value):
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заявки"
    headers = [
        "№ заявки", "Клиент", "Телефон", "Тип техники", "Описание проблемы",
        "Статус", "Дата создания", "Аккаунт клиента"
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="17233C")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30

    for row in rows:
        sheet.append([
            row["id"],
            safe_excel_value(row["customer_name"]),
            safe_excel_value(row["phone"]),
            safe_excel_value(row["device_type"]),
            safe_excel_value(row["problem_description"]),
            safe_excel_value(row["status"]),
            safe_excel_value(row["created_at"]),
            safe_excel_value(row["account_email"] or "Гостевая заявка"),
        ])

    widths = {"A": 12, "B": 25, "C": 22, "D": 22, "E": 55, "F": 20, "G": 28, "H": 32}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=(cell.column == 5))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    summary = workbook.create_sheet("Сводка")
    summary.append(["Показатель", "Значение"])
    summary.append(["Дата формирования", datetime.now().strftime("%d.%m.%Y %H:%M")])
    summary.append(["Заявок в отчёте", len(rows)])
    summary.append(["Всего зарегистрировано аккаунтов", user_count])
    summary.append(["Применённый фильтр статуса", selected_status])
    summary.append(["Поисковый запрос", search or "Не задан"])
    summary.append([])
    summary.append(["Статус", "Количество заявок"])
    for cell in (summary[1] + summary[8]):
        cell.fill = header_fill
        cell.font = header_font
    status_counts = {status: 0 for status in allowed_statuses}
    for item in rows:
        if item["status"] in status_counts:
            status_counts[item["status"]] += 1
    for status in allowed_statuses:
        summary.append([status, status_counts.get(status, 0)])
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 26
    summary.freeze_panes = "A2"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    suffix = datetime.now().strftime("%Y%m%d_%H%M%S")
    return send_file(
        output,
        as_attachment=True,
        download_name=f"computer_soft_requests_{suffix}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/admin/users")
@admin_required
def admin_users():
    """Показывает зарегистрированных пользователей и позволяет искать их."""
    search = (request.args.get("q") or "").strip()[:100]
    selected_role = (request.args.get("role") or "all").strip()
    if selected_role not in {"all", "user", "admin"}:
        selected_role = "all"

    conditions = []
    params = []
    if selected_role != "all":
        conditions.append("role = ?")
        params.append(selected_role)
    if search:
        conditions.append("(full_name LIKE ? OR email LIKE ? OR CAST(id AS TEXT) LIKE ?)")
        value = f"%{search}%"
        params.extend([value, value, value])

    sql = "SELECT id, full_name, email, created_at, role FROM users"
    if conditions:
        sql += " WHERE " + " AND ".join(conditions)
    sql += " ORDER BY id DESC LIMIT 500"

    with get_db_connection() as connection:
        users_list = connection.execute(sql, params).fetchall()
        total_users = connection.execute("SELECT COUNT(*) FROM users").fetchone()[0]
        total_admins = connection.execute("SELECT COUNT(*) FROM users WHERE role = 'admin'").fetchone()[0]

    return render_template(
        "admin_users.html",
        users_list=users_list,
        search=search,
        selected_role=selected_role,
        total_users=total_users,
        total_admins=total_admins,
        visible_count=len(users_list),
        current_admin_id=g.user["id"],
    )


@app.post("/admin/users/<int:user_id>/role")
@admin_required
def admin_update_user_role(user_id: int):
    """Меняет роль аккаунта, не позволяя администратору снять роль с себя."""
    new_role = (request.form.get("role") or "").strip()
    search = (request.form.get("q") or "").strip()[:100]
    selected_role = (request.form.get("filter_role") or "all").strip()
    if selected_role not in {"all", "user", "admin"}:
        selected_role = "all"

    if new_role not in {"user", "admin"}:
        flash("Выбрана неизвестная роль.", "error")
        return redirect(url_for("admin_users", q=search, role=selected_role))
    if user_id == g.user["id"] and new_role != "admin":
        flash("Нельзя снять роль администратора со своего аккаунта.", "error")
        return redirect(url_for("admin_users", q=search, role=selected_role))

    with get_db_connection() as connection:
        target = connection.execute(
            "SELECT id, full_name, email, role FROM users WHERE id = ?", (user_id,)
        ).fetchone()
        if target is not None:
            connection.execute("UPDATE users SET role = ? WHERE id = ?", (new_role, user_id))

    if target is None:
        flash("Пользователь не найден.", "error")
    elif target["role"] == new_role:
        flash("У пользователя уже установлена эта роль.", "success")
    else:
        role_label = "администратора" if new_role == "admin" else "пользователя"
        flash(f"Для аккаунта {target['email']} установлена роль: {role_label}.", "success")
    return redirect(url_for("admin_users", q=search, role=selected_role))


@app.errorhandler(413)
def too_large(_error):
    flash("Отправленные данные слишком большие.", "error")
    return redirect(url_for("index") + "#contact"), 413


if __name__ == "__main__":
    init_db()
    print("Компьютер-Софт: сайт запускается на http://127.0.0.1:5000")
    app.run(debug=True)
